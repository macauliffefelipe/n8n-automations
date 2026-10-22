#!/usr/bin/env python3
"""Genera el informe 'Monitoreo de medios' (formato vertical, 5 paginas aprox.) desde un JSON.
Uso: python3 gen_monitoreo.py datos.json salida.pptx
Si los datos son invalidos imprime la lista de errores y sale sin generar nada."""
import sys, json, re, math, io, base64
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.dml.color import RGBColor
from gen_informe import pct

W, H = 5.4, 9.6
ORANGE, DARK, GREY, BG = "FF6A13", "222222", "6B6B6B", "F4F4F4"
SERIF, SANS = "Georgia", "Calibri"
BAR_COLORS = ["FF6A13", "FFC000", "6D6D6D", "9E9E9E"]
SENT_COLORS = {"positivo": "34A853", "negativo": "E53935", "neutral": "C8C8C8"}
LIM = {"titular": 140, "texto": 420, "tema": 60, "hl_titular": 140}

def rgb(h): return RGBColor.from_string(h)

def validate(d):
    e = []
    def need(k, t=None):
        if k not in d or d[k] in (None, ""): e.append(f"falta '{k}'"); return False
        return True
    for k in ("cliente", "cliente_corto", "fecha"): need(k)
    if need("corte") and not re.fullmatch(r"\d{1,2}:\d{2}", str(d["corte"])): e.append("'corte' debe ser HH:MM")
    if need("menciones") and not (isinstance(d["menciones"], int) and d["menciones"] >= 0): e.append("'menciones' debe ser entero >= 0")
    if need("alcance") and not isinstance(d["alcance"], (int, str)): e.append("'alcance' debe ser numero o texto")
    t = d.get("temas")
    if not t: e.append("falta 'temas'")
    else:
        if len(t) > 4: e.append("maximo 4 temas")
        for i, x in enumerate(t):
            if not (isinstance(x, list) and len(x) == 2 and isinstance(x[1], int) and x[1] >= 0): e.append(f"tema {i+1}: formato [nombre, cantidad]"); continue
            if len(x[0]) > LIM["tema"]: e.append(f"tema {i+1}: {len(x[0])} caracteres (maximo {LIM['tema']})")
    s = d.get("sentiment")
    if not isinstance(s, dict) or any(k not in s or not isinstance(s[k], int) for k in ("positivo", "neutral", "negativo")): e.append("'sentiment' necesita positivo, neutral y negativo (enteros)")
    elif sum(s.values()) == 0: e.append("'sentiment' suma 0")
    c = d.get("tarjetas")
    if not c: e.append("falta 'tarjetas'")
    else:
        if len(c) > 6: e.append("maximo 6 tarjetas")
        for i, x in enumerate(c):
            if not (isinstance(x, dict) and x.get("titular") and x.get("texto")): e.append(f"tarjeta {i+1}: necesita titular y texto"); continue
            if len(x["titular"]) > LIM["titular"]: e.append(f"tarjeta {i+1}: titular {len(x['titular'])} caracteres (maximo {LIM['titular']})")
            if len(x["texto"]) > LIM["texto"]: e.append(f"tarjeta {i+1}: texto {len(x['texto'])} caracteres (maximo {LIM['texto']})")
    if d.get("fuente_nota") and len(d["fuente_nota"]) > 220: e.append("'fuente_nota' supera 220 caracteres")
    h = d.get("highlights", [])
    if len(h) > 6: e.append("maximo 6 highlights")
    for i, x in enumerate(h):
        if not (isinstance(x, dict) and x.get("medio") and x.get("titular")): e.append(f"highlight {i+1}: necesita medio y titular"); continue
        if len(x["titular"]) > LIM["hl_titular"]: e.append(f"highlight {i+1}: titular {len(x['titular'])} caracteres (maximo {LIM['hl_titular']})")
    return e

def fmt_alcance(a):
    if isinstance(a, str): return a
    if a >= 1_000_000: return f"{a/1e6:.1f}M".replace(".0M", "M")
    if a >= 1_000: return f"{round(a/1e3)}K"
    return str(a)

def rect(sl, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, line=None):
    s = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    if line: s.line.color.rgb = rgb(line); s.line.width = Pt(0.75)
    else: s.line.fill.background()
    s.shadow.inherit = False
    return s

def text(sl, x, y, w, h, t, size, color=DARK, font=SANS, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = align
    if spacing: p.line_spacing = spacing
    r = p.add_run(); r.text = t; f = r.font; f.size = Pt(size); f.name = font; f.bold = bold; f.color.rgb = rgb(color)
    return tb

def new_page(prs, bg=BG):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    sl.background.fill.solid(); sl.background.fill.fore_color.rgb = rgb(bg)
    return sl

def header(sl, d):
    rect(sl, 0, 0, W, 1.15, "FFFFFF")
    text(sl, 0.45, 0.28, 3.8, 0.4, f"{d['cliente_corto']} | Monitoreo", 20, DARK, SERIF, True)
    text(sl, 0.45, 0.72, 3.8, 0.25, f"{d['fecha']} | Corte {d['corte']}", 9, DARK)
    c = rect(sl, W-1.2, 0.2, 0.75, 0.75, ORANGE, MSO_SHAPE.OVAL)
    text(sl, W-1.2, 0.2, 0.75, 0.75, "D.", 20, "FFFFFF", SERIF, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

def section(sl, y, title):
    rect(sl, 0.45, y, 0.07, 0.45, ORANGE)
    text(sl, 0.7, y, 4.4, 0.5, title, 14, DARK, SERIF, True, anchor=MSO_ANCHOR.MIDDLE)

def cover(prs, d):
    sl = new_page(prs, ORANGE)
    rect(sl, 0, H*0.62, W, H*0.38, "E85A0A")
    text(sl, 0.4, 3.3, W-0.8, 1.6, "Conversación digital", 34, "FFFFFF", SERIF, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    pill = rect(sl, 1.35, 5.05, 2.7, 0.5, "FFFFFF", MSO_SHAPE.ROUNDED_RECTANGLE); pill.adjustments[0] = 0.5
    text(sl, 1.35, 5.05, 2.7, 0.5, "BY YOUR AGENCY", 11, ORANGE, SANS, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    text(sl, 0.4, 7.4, W-0.8, 0.7, d["cliente"], 22, "FFFFFF", SANS, False, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

def style_chart_text(chart, size=8):
    chart.font.size = Pt(size); chart.font.name = SANS; chart.font.color.rgb = rgb(DARK)

def overview(prs, d):
    sl = new_page(prs); header(sl, d)
    section(sl, 1.4, f"Conversación digital alrededor de {d['cliente_corto']}")
    names = [t[0] for t in d["temas"]]; counts = [t[1] for t in d["temas"]]; share = pct(counts)
    cd = CategoryChartData(); cd.categories = names; cd.add_series("Share", share)
    gf = sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.45), Inches(2.0), Inches(W-0.9), Inches(2.7), cd)
    ch = gf.chart; style_chart_text(ch, 7); ch.has_legend = False; ch.has_title = False
    ch.value_axis.visible = False; ch.value_axis.has_major_gridlines = False; ch.value_axis.maximum_scale = max(share) * 1.2
    ch.category_axis.tick_labels.font.size = Pt(7); ch.category_axis.tick_labels.font.bold = True
    ch.category_axis.format.line.color.rgb = rgb("BBBBBB")
    pl = ch.plots[0]; pl.gap_width = 35; pl.has_data_labels = True
    dl = pl.data_labels; dl.number_format = '0"%"'; dl.number_format_is_linked = False; dl.position = XL_LABEL_POSITION.OUTSIDE_END; dl.font.size = Pt(9); dl.font.bold = True; dl.show_value = True
    for i in range(len(names)):
        pt = pl.series[0].points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(BAR_COLORS[i % 4])
    section(sl, 4.9, "Volumen")
    rect(sl, 0.7, 5.55, W-1.4, 0.95, "FFFFFF", MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.1
    half = (W-1.4)/2
    text(sl, 0.7, 5.62, half, 0.5, f"{d['menciones']:,}".replace(",", "."), 24, ORANGE, SERIF, True, PP_ALIGN.CENTER)
    text(sl, 0.7, 6.12, half, 0.3, "menciones", 12, DARK, SANS, False, PP_ALIGN.CENTER)
    rect(sl, W/2, 5.7, 0.01, 0.65, "CCCCCC")
    text(sl, W/2, 5.62, half, 0.5, fmt_alcance(d["alcance"]), 24, ORANGE, SERIF, True, PP_ALIGN.CENTER)
    text(sl, W/2, 6.12, half, 0.3, "alcance", 12, DARK, SANS, False, PP_ALIGN.CENTER)
    section(sl, 6.75, "Sentiment")
    s = d["sentiment"]; order = ["positivo", "negativo", "neutral"]; sp = pct([s[k] for k in order])
    cd2 = CategoryChartData(); cd2.categories = [k.capitalize() for k in order]; cd2.add_series("Sentiment", sp)
    gf2 = sl.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(0.7), Inches(7.3), Inches(W-1.4), Inches(1.95), cd2)
    c2 = gf2.chart; style_chart_text(c2, 9); c2.has_title = False; c2.has_legend = True
    c2.legend.position = XL_LEGEND_POSITION.RIGHT; c2.legend.include_in_layout = False; c2.legend.font.size = Pt(9)
    p2 = c2.plots[0]; p2.has_data_labels = True
    d2 = p2.data_labels; d2.number_format = '0"%"'; d2.number_format_is_linked = False; d2.font.size = Pt(9); d2.font.bold = True; d2.font.color.rgb = rgb("FFFFFF"); d2.show_value = True; d2.show_category_name = False; d2.show_percentage = False
    for i, k in enumerate(order):
        pt = p2.series[0].points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(SENT_COLORS[k])
    if d.get("fuente_nota"):
        text(sl, 0.45, 9.27, W-0.9, 0.3, d["fuente_nota"], 6.5, GREY, SANS, align=PP_ALIGN.CENTER)
    hs = p2._element.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}holeSize')
    if hs is not None: hs.set("val", "55")

def cards(prs, d):
    per, y0, ch, gap = 3, 1.4, 2.5, 0.2
    for pi in range(0, len(d["tarjetas"]), per):
        sl = new_page(prs); header(sl, d)
        for j, c in enumerate(d["tarjetas"][pi:pi+per]):
            y = y0 + j*(ch+gap)
            rect(sl, 0.4, y, W-0.8, ch, "FFFFFF", line="E6E6E6"); rect(sl, 0.4, y, 0.07, ch, ORANGE)
            rect(sl, 0.7, y+0.25, 0.5, 0.5, ORANGE, MSO_SHAPE.OVAL)
            text(sl, 0.7, y+0.25, 0.5, 0.5, str(pi+j+1), 14, "FFFFFF", SERIF, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
            # El texto arranca debajo del titular segun cuantas lineas ocupe (aprox. 30 caracteres por linea a 11 pt)
            lineas = max(2, math.ceil(len(c["titular"]) / 30))
            ty = max(y + 1.1, y + 0.2 + lineas * 0.2 + 0.15)
            text(sl, 1.4, y+0.2, W-2.0, ty - y - 0.25, c["titular"], 11, DARK, SERIF, True)
            text(sl, 0.7, ty, W-1.5, y + ch - ty - 0.1, c["texto"], 10, "333333", SANS, spacing=1.05)

def b64img(s):
    return io.BytesIO(base64.b64decode(s))

def highlights(prs, d):
    hl = d.get("highlights", [])
    if not hl: return
    sl = new_page(prs); header(sl, d); section(sl, 1.4, "Highlights")
    cw, chh, gx, gy, x0, y0 = 2.3, 2.45, 0.2, 0.15, 0.4, 2.1
    for i, x in enumerate(hl):
        col, row = i % 2, i // 2
        X, Y = x0 + col*(cw+gx), y0 + row*(chh+gy)
        rect(sl, X, Y, cw, chh, "FFFFFF", line="E6E6E6"); rect(sl, X, Y, 0.06, chh, ORANGE)
        ty = Y + 0.15
        if x.get("captura_b64"):
            try:
                sl.shapes.add_picture(b64img(x["captura_b64"]), Inches(X+0.2), Inches(Y+0.15), width=Inches(cw-0.35), height=Inches(1.25)); ty = Y + 1.5
            except Exception: pass
        meta = x["medio"] + (f"  ·  {x['hora']}" if x.get("hora") else "")
        text(sl, X+0.2, ty, cw-0.35, 0.25, meta, 8, ORANGE, SANS, True)
        text(sl, X+0.2, ty+0.28, cw-0.35, Y+chh-ty-0.35, x["titular"], 10, DARK, SERIF, True)

def banner(sl):
    rect(sl, 0, 0, W, 0.24, "C62828")
    text(sl, 0, 0, W, 0.24, "DATOS DE EJEMPLO, NO ES UN REPORTE REAL", 8, "FFFFFF", SANS, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

def build(d, out):
    prs = Presentation(); prs.slide_width = Inches(W); prs.slide_height = Inches(H)
    cover(prs, d); overview(prs, d); cards(prs, d); highlights(prs, d)
    if d.get("ejemplo"):
        for sl in prs.slides: banner(sl)
    prs.save(out)

if __name__ == "__main__":
    if len(sys.argv) != 3: raise SystemExit("Uso: python3 gen_monitoreo.py datos.json salida.pptx")
    d = json.load(open(sys.argv[1], encoding="utf-8")); errs = validate(d)
    if errs:
        print("DATOS INVALIDOS:"); [print(" -", x) for x in errs]; raise SystemExit(1)
    build(d, sys.argv[2]); print("OK", sys.argv[2])
