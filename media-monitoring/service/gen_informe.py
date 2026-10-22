#!/usr/bin/env python3
"""Genera el informe AM/PM de 7 laminas desde un JSON y la plantilla. Uso: gen_informe.py datos.json plantilla.pptx salida.pptx"""
import sys, json, copy, math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.dml.color import RGBColor

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
LIMITS = {"lectura_rapida": 220, "sentiment_texto": 160, "que_esta_pasando": 330, "movimiento_principal": 110, "riesgo_oportunidad": 120, "que_seguir": 110}

def pct(counts):
    """Porcentajes enteros que suman 100 (resto mayor)."""
    tot = sum(counts)
    if tot == 0: return [0]*len(counts)
    raw = [100*c/tot for c in counts]; base = [int(r) for r in raw]
    for i in sorted(range(len(raw)), key=lambda i: raw[i]-base[i], reverse=True)[:100-sum(base)]: base[i] += 1
    return base

def validate(d):
    errs = []
    for k in ("cliente","turno","fecha","corte","objetivo","menciones","alcance","interacciones","lectura_rapida","sentiment","sentiment_texto","temas","interlocutores","analisis","highlights","ficha"):
        if k not in d or d[k] in ("", None, []): errs.append(f"falta '{k}'")
    if d.get("turno") and d["turno"] not in ("AM","PM"): errs.append("turno debe ser AM o PM")
    s = d.get("sentiment") or {}
    if s and set(s) != {"positivo","neutral","negativo"} or any(not isinstance(v,int) or v<0 for v in s.values()): errs.append("sentiment: enteros positivo/neutral/negativo")
    if len(d.get("temas") or []) > 4: errs.append("maximo 4 temas")
    if len(d.get("interlocutores") or []) > 4: errs.append("maximo 4 interlocutores")
    if len(d.get("highlights") or []) > 6: errs.append("maximo 6 highlights")
    for k, lim in LIMITS.items():
        v = d.get(k) or (d.get("analisis") or {}).get(k, "")
        if len(v) > lim: errs.append(f"'{k}' tiene {len(v)} caracteres (maximo {lim})")
    return errs

def find(slide, prefix, exact=False):
    return [s for s in slide.shapes if s.has_text_frame and ((s.text_frame.text.strip()==prefix) if exact else s.text_frame.text.strip().startswith(prefix))]

def set_text(sh, text, size=None, **geo):
    tf = sh.text_frame; para = tf.paragraphs[0]
    if not para.runs: para.add_run().text = text
    else:
        para.runs[0].text = text
        for r in para.runs[1:]: r._r.getparent().remove(r._r)
    for extra in tf.paragraphs[1:]: extra._p.getparent().remove(extra._p)
    if size:
        for r in para.runs: r.font.size = Pt(size)
    for k in ("left","top","width","height"):
        if k in geo and geo[k] is not None: setattr(sh, k, Inches(geo[k]))
    tf.word_wrap = True

def build(d, tpl, out):
    p = Presentation(tpl); S = list(p.slides)
    # 1 portada
    set_text(find(S[0], "[CLIENTE]")[0], d["cliente"]); set_text(find(S[0], "[AM")[0], d["turno"]); set_text(find(S[0], "[DD")[0], d["fecha"])
    # 2 estadisticas
    s = S[1]; vals = sorted(find(s, "[X]", True), key=lambda x: x.top)
    for sh, k in zip(vals, ("menciones","alcance","interacciones")): set_text(sh, str(d[k]))
    nota = lambda k: "" if str(d[k]) != "n/d" else " (sin monitoreo)"
    set_text(find(s, "Menciones del período")[0], "Menciones" + (" (muestra manual)" if d.get("muestra_manual", True) else " del período"))
    set_text(find(s, "Alcance del período")[0], "Alcance" + nota("alcance")); set_text(find(s, "Interacciones")[0], "Interacciones" + nota("interacciones"))
    set_text(find(s, "[Qué movió")[0], d["lectura_rapida"], top=5.2, height=0.75)
    # 3 sentiment
    s = S[2]; sv = d["sentiment"]; counts = [sv["positivo"], sv["neutral"], sv["negativo"]]; pc = pct(counts)
    box = [x for x in s.shapes if x.name.startswith("Rounded Rectangle")][0]
    for t in ("GRÁFICO SENTIMENT", "Positivo · Neutral"):
        for sh in find(s, t): sh._element.getparent().remove(sh._element)
    cd = CategoryChartData(); cd.categories = ["Positivo","Neutral","Negativo"]; cd.add_series("Sentiment", counts)
    gf = s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, box.left+Inches(0.55), box.top+Inches(0.08), box.width-Inches(1.1), box.height-Inches(0.16), cd)
    gf.chart.has_legend = False; gf.chart.has_title = False; gf.chart.plots[0].has_data_labels = False
    for i, c in enumerate((RGBColor(0x4C,0xAF,0x6A), RGBColor(0x9E,0x9E,0x9E), RGBColor(0xE8,0x45,0x3C))):
        pt = gf.chart.plots[0].series[0].points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = c
    for sh, (n, v) in zip((find(s,"POSITIVO")[0], find(s,"NEUTRAL")[0], find(s,"NEGATIVO")[0]), zip(("POSITIVO","NEUTRAL","NEGATIVO"), pc)): set_text(sh, f"{n} {v}%")
    for sh, c in zip(sorted(find(s,"[X]",True), key=lambda x: x.left), counts): set_text(sh, str(c))
    set_text(find(s, "[Explicación")[0], d["sentiment_texto"], top=3.98, height=0.55)
    # 4 temas e interlocutores
    s = S[3]; temas = d["temas"]; tp = pct([c for _, c in temas])
    names = sorted(find(s,"[Tema / narrativa]",True), key=lambda x: x.top); pcts = sorted(find(s,"[XX%]",True), key=lambda x: x.top)
    for i, (sh, sp) in enumerate(zip(names, pcts)):
        if i < len(temas): set_text(sh, temas[i][0], width=1.95); set_text(sp, f"{tp[i]}%")
        else: set_text(sh, ""); set_text(sp, "")
    wn = sorted(find(s,"[Perfil / medio]",True), key=lambda x: x.top); wf = sorted(find(s,"[Link-seguidores]",True), key=lambda x: x.top)
    for i, (a, b) in enumerate(zip(wn, wf)):
        if i < len(d["interlocutores"]): set_text(a, d["interlocutores"][i][0]); set_text(b, d["interlocutores"][i][1])
        else: set_text(a, ""); set_text(b, "")
    # 5 analisis
    s = S[4]; an = d["analisis"]
    set_text(find(s, "[Síntesis")[0], an["que_esta_pasando"], height=1.2)
    cc = sorted(find(s,"[Conclusión breve]",True), key=lambda x: x.top)
    for sh, k in zip(cc, ("movimiento_principal","riesgo_oportunidad","que_seguir")): set_text(sh, an[k])
    # 6 highlights
    s = S[5]; set_text(find(s,"HIGLIGHTS")[0], "HIGHLIGHTS")
    t = find(s, "X", True)
    if t: set_text(t[0], d.get("highlights_titulo", "Medios"))
    rects = [x for x in s.shapes if x.name.startswith("Rounded Rectangle")]
    caps = sorted([sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip().startswith("CAPTURA")], key=lambda x: int(x.text_frame.text.split()[1]))
    for i, sh in enumerate(caps):
        cx, cy = sh.left+sh.width//2, sh.top+sh.height//2
        r = [x for x in rects if x.left<=cx<=x.left+x.width and x.top<=cy<=x.top+x.height][0]
        if i >= len(d["highlights"]): set_text(sh, ""); continue
        medio, hora, titular = d["highlights"][i]
        set_text(sh, f"{medio} · {hora}", size=6.5, left=r.left/914400+0.06, top=r.top/914400+0.06, width=r.width/914400-0.12, height=0.2)
        sh.text_frame.vertical_anchor = MSO_ANCHOR.TOP
        for rr in sh.text_frame.paragraphs[0].runs: rr.font.bold = True
        para = sh.text_frame.add_paragraph(); run = para.add_run(); run.text = titular
        src = sh.text_frame.paragraphs[0].runs[0]._r.find(A+"rPr"); old = run._r.find(A+"rPr")
        if old is not None: run._r.remove(old)
        run._r.insert(0, copy.deepcopy(src)); run.font.bold = False
        sh.height = r.height - Inches(0.12)
        for pp in sh.text_frame.paragraphs: pp.alignment = PP_ALIGN.LEFT
    # 7 ficha
    s = S[6]
    set_text(find(s, "[AM / PM")[0], d["corte"]); set_text(find(s, "[Qué se releva")[0], d["objetivo"], height=0.5)
    lab = [x for x in s.shapes if x.has_text_frame and x.text_frame.text.strip()=="OBJETIVO"][0]
    val = [x for x in s.shapes if x.has_text_frame and x.text_frame.text.strip()==d["objetivo"]][0]
    top = val.top/914400 + 0.62
    for ltxt, vtxt in d["ficha"]:
        h = 0.15*math.ceil(len(vtxt)/50) + 0.1
        nl, nv = copy.deepcopy(lab._element), copy.deepcopy(val._element)
        s.shapes._spTree.append(nl); s.shapes._spTree.append(nv)
        set_text(s.shapes[-2], ltxt, top=top, width=2.95); set_text(s.shapes[-1], vtxt, top=top+0.25, height=h)
        top += 0.25 + h + 0.1
    if top > p.slide_height/914400 - 0.1: raise SystemExit(f"ERROR: la ficha tecnica no entra en la lamina (llega a {top:.2f} in)")
    p.save(out)

if __name__ == "__main__":
    d = json.load(open(sys.argv[1], encoding="utf-8"))
    errs = validate(d)
    if errs: raise SystemExit("DATOS INVALIDOS:\n- " + "\n- ".join(errs))
    build(d, sys.argv[2], sys.argv[3]); print("OK", sys.argv[3])
